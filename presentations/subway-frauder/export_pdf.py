from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader


ROOT = Path(__file__).resolve().parent
FONT_DIR = Path('/System/Library/Fonts/Supplemental')
for label, filename in [('regular', 'Arial.ttf'), ('bold', 'Arial Bold.ttf'),
                        ('italic', 'Arial Italic.ttf'), ('bolditalic', 'Arial Bold Italic.ttf')]:
    pdfmetrics.registerFont(TTFont(label, str(FONT_DIR / filename)))

prs = Presentation(ROOT / 'Subway-Frauder-Presentation.pptx')
W, H = prs.slide_width.pt, prs.slide_height.pt
c = canvas.Canvas(str(ROOT / 'Subway-Frauder-Presentation.pdf'), pagesize=(W, H), pageCompression=1)
c.setTitle('Subway Frauder — Bouge pour jouer')
c.setAuthor('Ali BEN YEZZA · Noé WALES')
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def get_fill(fill):
    if fill.type != MSO_FILL_TYPE.SOLID:
        return None, 1
    color = HexColor('#' + str(fill.fore_color.rgb))
    alpha = fill._xPr.find('.//' + A + 'alpha')
    return color, int(alpha.get('val')) / 100000 if alpha is not None else 1


def arrow_path(x, y, w, h, direction):
    points = [(0, .25), (.55, .25), (.55, 0), (1, .5), (.55, 1), (.55, .75), (0, .75)]
    if direction == MSO_AUTO_SHAPE_TYPE.LEFT_ARROW:
        points = [(1 - px, py) for px, py in points]
    elif direction == MSO_AUTO_SHAPE_TYPE.UP_ARROW:
        points = [(py, px) for px, py in points]
    elif direction == MSO_AUTO_SHAPE_TYPE.DOWN_ARROW:
        points = [(py, 1 - px) for px, py in points]
    p = c.beginPath()
    p.moveTo(x + points[0][0] * w, y + points[0][1] * h)
    for px, py in points[1:]:
        p.lineTo(x + px * w, y + py * h)
    p.close()
    return p


for slide in prs.slides:
    background, _ = get_fill(slide.background.fill)
    c.setFillColor(background or HexColor('#172B40'))
    c.rect(0, 0, W, H, stroke=0, fill=1)
    for s in slide.shapes:
        x, top, w, h = s.left.pt, s.top.pt, s.width.pt, s.height.pt
        y = H - top - h
        c.saveState()
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
            p = c.beginPath()
            geom = s._element.spPr.find(A + 'prstGeom')
            if geom is not None and geom.get('prst') == 'ellipse':
                p.ellipse(x, y, w, h)
            else:
                p.rect(x, y, w, h)
            c.clipPath(p, stroke=0)
            dw = w / (1 - s.crop_left - s.crop_right)
            dh = h / (1 - s.crop_top - s.crop_bottom)
            c.drawImage(ImageReader(BytesIO(s.image.blob)), x - s.crop_left * dw,
                        y - s.crop_bottom * dh, dw, dh, mask='auto')
        elif s.shape_type == MSO_SHAPE_TYPE.LINE:
            color, _ = get_fill(s.line.fill)
            c.setStrokeColor(color or HexColor('#496176'))
            c.setLineWidth(s.line.width.pt)
            c.line(x, H - top, x + w, H - top - h)
        elif s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            fill, alpha = get_fill(s.fill)
            stroke, _ = get_fill(s.line.fill)
            if fill:
                c.setFillColor(fill)
                c.setFillAlpha(alpha)
            if stroke:
                c.setStrokeColor(stroke)
                c.setLineWidth(s.line.width.pt)
            fill_flag, stroke_flag = int(fill is not None), int(stroke is not None)
            if s.auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL:
                c.ellipse(x, y, x + w, y + h, stroke=stroke_flag, fill=fill_flag)
            elif s.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
                c.roundRect(x, y, w, h, min(w, h) * .09, stroke=stroke_flag, fill=fill_flag)
            elif s.auto_shape_type in [MSO_AUTO_SHAPE_TYPE.LEFT_ARROW, MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
                                        MSO_AUTO_SHAPE_TYPE.UP_ARROW, MSO_AUTO_SHAPE_TYPE.DOWN_ARROW]:
                c.drawPath(arrow_path(x, y, w, h, s.auto_shape_type), stroke=stroke_flag, fill=fill_flag)
            else:
                c.rect(x, y, w, h, stroke=stroke_flag, fill=fill_flag)
        elif s.has_text_frame:
            offset = 0
            for paragraph in s.text_frame.paragraphs:
                if not paragraph.runs:
                    continue
                run = paragraph.runs[0]
                size = run.font.size.pt
                family = ('bold' if run.font.bold else '') + ('italic' if run.font.italic else '') or 'regular'
                color = HexColor('#' + str(run.font.color.rgb))
                tracking = float(run._r.get_or_add_rPr().get('spc', '0')) / 100
                text_width = pdfmetrics.stringWidth(run.text, family, size) + tracking * max(0, len(run.text) - 1)
                tx = x
                if paragraph.alignment == PP_ALIGN.CENTER:
                    tx += (w - text_width) / 2
                elif paragraph.alignment == PP_ALIGN.RIGHT:
                    tx += w - text_width
                baseline = H - top - size * .94 - offset
                t = c.beginText(tx, baseline)
                t.setFont(family, size)
                t.setFillColor(color)
                t.setCharSpace(tracking)
                outline = run._r.get_or_add_rPr().find(A + 'ln')
                if outline is not None:
                    t.setTextRenderMode(2)
                    c.setStrokeColor(HexColor('#0C192A'))
                    c.setLineWidth(.7)
                t.textOut(run.text)
                c.drawText(t)
                if run.hyperlink.address:
                    c.linkURL(run.hyperlink.address, (x, y, x + w, y + h), relative=0)
                offset += paragraph.line_spacing.pt if hasattr(paragraph.line_spacing, 'pt') else size * 1.12
        c.restoreState()
    c.showPage()
c.save()
print('PDF exporté :', ROOT / 'Subway-Frauder-Presentation.pdf')
