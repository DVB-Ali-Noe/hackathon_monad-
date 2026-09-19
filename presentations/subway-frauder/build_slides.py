from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
W, H = 960, 540
NAVY = '172B40'
DEEP = '0C192A'
PANEL = '20364C'
WHITE = 'F6F7F8'
YELLOW = 'F9D85C'
PURPLE = '9167ED'
CYAN = 'ACD4E4'
MUTED = 'AEC0CE'
LINE = '496176'

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = 'Subway Frauder — Bouge pour jouer'
prs.core_properties.subject = 'Monad Blitz Paris 2026'
prs.core_properties.author = 'Ali BEN YEZZA · Noé WALES'
prs.core_properties.keywords = 'Subway Frauder, Monad, DeVinci Blockchain, webcam'


def rgb(value):
    return RGBColor.from_string(value)


def opacity(color_element, value):
    alpha = OxmlElement('a:alpha')
    alpha.set('val', str(round(value * 100000)))
    color_element.append(alpha)


def shape(slide, x, y, w, h, fill=None, stroke=None, line_width=1, radius=False,
          kind=None, alpha=1):
    kind = kind or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if radius:
        s.adjustments[0] = 0.09
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if alpha != 1:
            opacity(s.fill._xPr.solidFill.srgbClr, alpha)
    else:
        s.fill.background()
    if stroke:
        s.line.color.rgb = rgb(stroke)
        s.line.width = Pt(line_width)
    else:
        s.line.fill.background()
    return s


def text(slide, value, x, y, w, h, size=20, color=WHITE, bold=False,
         italic=False, align='left', font='Arial', spacing=1.12, tracking=None,
         outline=False):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    box.text_frame.clear()
    box.text_frame.word_wrap = False
    box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    box.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    box.text_frame.margin_left = box.text_frame.margin_right = 0
    box.text_frame.margin_top = box.text_frame.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size * spacing)
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold, r.font.italic = bold, italic
        r.font.color.rgb = rgb(color)
        rpr = r._r.get_or_add_rPr()
        rpr.set('lang', 'fr-FR')
        if tracking is not None:
            rpr.set('spc', str(int(tracking * 100)))
        if outline:
            border = OxmlElement('a:ln')
            border.set('w', '12700')
            fill = OxmlElement('a:solidFill')
            c = OxmlElement('a:srgbClr')
            c.set('val', DEEP)
            fill.append(c)
            border.append(fill)
            rpr.insert(0, border)
    return box


def line(slide, x1, y1, x2, y2, color=LINE, width=1):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    return s


def picture(slide, name, x, y, w, h, crop=None, circle=False):
    p = slide.shapes.add_picture(str(ASSETS / name), Pt(x), Pt(y), Pt(w), Pt(h))
    if crop:
        left, top, right, bottom = crop
        p.crop_left, p.crop_top = left, top
        p.crop_right, p.crop_bottom = 1 - right, 1 - bottom
    if circle:
        geometry = p._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
        geometry.set('prst', 'ellipse')
    return p


def chip(slide, value, x, y, w, fill, color=WHITE, size=10):
    shape(slide, x, y, w, 25, fill=fill, radius=True)
    text(slide, value, x, y + 6, w, 16, size, color, True, align='center', tracking=0.7)


def footer(slide, number, color=MUTED):
    text(slide, 'SUBWAY FRAUDER  /  MONAD BLITZ PARIS', 52, 514, 700, 14, 8, color, tracking=1)
    text(slide, f'{number:02d} / 05', 834, 514, 74, 14, 9, color, align='right')


def base(number, label, background=NAVY):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(background)
    text(s, 'SUBWAY', 52, 25, 66, 15, 11, WHITE, True)
    text(s, 'FRAUDER', 120, 25, 80, 15, 11, YELLOW, True)
    text(s, label, 535, 27, 373, 15, 9, CYAN, True, align='right', tracking=1.1)
    line(s, 52, 55, 908, 55, LINE, 0.6)
    footer(s, number)
    return s


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


# 1 — Une couverture illustrée, sans interface de navigateur.
s = prs.slides.add_slide(prs.slide_layouts[6])
picture(s, 'cover-paris.png', 0, 0, W, H)
shape(s, 0, 0, W, H, DEEP, alpha=0.14)
for i in range(72):
    x = i * 9
    shade = 0.87 * (1 - i / 71) ** 0.85
    shape(s, x, 0, 9.1, H, DEEP, alpha=shade)
chip(s, 'PARIS', 52, 37, 70, '3E9C80')
chip(s, 'ON MONAD', 131, 37, 116, '7952D4')
text(s, 'SUBWAY', 49, 115, 464, 94, 78, WHITE, True, True, outline=True)
text(s, 'FRAUDER', 48, 198, 470, 94, 78, YELLOW, True, True, outline=True)
text(s, 'BOUGE POUR JOUER.', 54, 306, 458, 35, 25, WHITE, True, True)
text(s, 'Un runner à la webcam.\nDes scores sur Monad.', 55, 354, 385, 59, 19, WHITE, spacing=1.35)
shape(s, 54, 452, 317, 34, YELLOW, radius=True)
text(s, 'Zéro manette. Un peu de cardio.', 63, 460, 298, 22, 15, DEEP, True, align='center')
footer(s, 1, 'D3DEE8')
notes(s, "20 s — Subway Frauder, c’est un runner qui transforme le corps en manette. On joue devant sa webcam : on esquive, on saute, on se baisse. Le parcours est parisien et les scores sont enregistrés sur Monad testnet. Accroche : ‘On voulait faire un jeu blockchain. On a aussi fait une séance de cardio.’")


# 2 — Les portraits fournis restent de vraies images, cadrées nativement.
s = base(2, 'DEVINCI BLOCKCHAIN')
text(s, 'QUI SOMMES-NOUS ?', 52, 82, 810, 65, 43, WHITE, True, True)
text(s, 'Deux développeurs. Une webcam. Beaucoup trop d’idées.', 54, 143, 850, 35, 19, MUTED)
for cx, color in [(270, YELLOW), (690, PURPLE)]:
    shape(s, cx - 105, 205, 210, 210, None, color, 3, kind=MSO_SHAPE.OVAL)
picture(s, 'equipe-reference.png', 169, 209, 202, 202, (0.202, 0.321, 0.445, 0.724), circle=True)
picture(s, 'equipe-reference.png', 589, 209, 202, 202, (0.556, 0.321, 0.799, 0.724), circle=True)
chip(s, 'JOUEUR 01', 218, 186, 104, YELLOW, DEEP)
chip(s, 'JOUEUR 02', 638, 186, 104, PURPLE)
text(s, 'Ali BEN YEZZA', 94, 436, 352, 40, 28, WHITE, True, align='center')
text(s, 'Noé WALES', 514, 436, 352, 40, 28, WHITE, True, align='center')
for i in range(16):
    shape(s, 464 + i * 22, 501 - i * 6.7, 3, 3, LINE, kind=MSO_SHAPE.OVAL)
notes(s, "20 s — Nous sommes Ali Ben Yezza et Noé Wales, de DeVinci Blockchain. Nous avons construit Subway Frauder pour Monad Blitz Paris. Deux développeurs, une webcam et l’envie de proposer une démo qu’on peut essayer immédiatement. Pas besoin d’inventer des rôles ou des biographies : présentez simplement chacun votre contribution à l’oral.")


# 3 — Des consignes de jeu lisibles à distance.
s = base(3, 'COMMENT ÇA SE JOUE')
text(s, 'LA MANETTE,', 52, 87, 630, 63, 48, WHITE, True, True)
text(s, 'C’EST TOI.', 50, 145, 630, 63, 48, YELLOW, True, True)
text(s, 'Une webcam.\nUn navigateur.\nEt des jambes.', 703, 97, 215, 104, 21, WHITE, spacing=1.35)
commands = [
    ('GAUCHE', 'Change de voie', MSO_SHAPE.LEFT_ARROW, YELLOW),
    ('DROITE', 'Évite les obstacles', MSO_SHAPE.RIGHT_ARROW, YELLOW),
    ('SAUT', 'Passe au-dessus', MSO_SHAPE.UP_ARROW, PURPLE),
    ('ACCROUPI', 'Passe en dessous', MSO_SHAPE.DOWN_ARROW, PURPLE),
]
for i, (title, caption, arrow, color) in enumerate(commands):
    x = 52 + 220 * i
    shape(s, x, 249, 196, 168, PANEL, LINE, 1.1, True)
    line(s, x + 17, 249, x + 58, 249, color, 2.6)
    shape(s, x + 73, 271, 50, 47, color, kind=arrow)
    text(s, title, x + 8, 338, 180, 28, 19, WHITE, True, align='center')
    text(s, caption, x + 8, 373, 180, 24, 13, MUTED, align='center')
text(s, 'Le bouton « sauter », c’est toi.', 52, 451, 603, 33, 24, YELLOW, True, True)
chip(s, 'VIDÉO TRAITÉE EN LOCAL', 677, 456, 232, '2B4559', CYAN, 9)
notes(s, "35 s — On choisit un pseudo, on active la caméra et on a cinq secondes pour se placer. Le corps contrôle gauche, droite, saut et accroupissement. Le parcours est généré à partir d’une graine aléatoire. La caméra et la simulation tournent dans le navigateur ; les images ne sont pas envoyées au serveur. Le clavier et le tactile restent disponibles pour essayer le jeu. Si possible, montrer les mouvements devant la salle.")


# 4 — La capture fournie sert de preuve visuelle, sans modifier ses chiffres.
s = base(4, 'LE RÉSULTAT DU RUN')
text(s, 'LE RUN FINIT. LA RIVALITÉ RESTE.', 51, 79, 860, 48, 35, WHITE, True, True)
text(s, 'Score, pièces, highlights locaux et classement.', 54, 130, 850, 24, 17, MUTED)
shape(s, 48, 160, 864, 338, DEEP, LINE, 1.2, True)
picture(s, 'resultats-reference.jpg', 54, 165, 852, 329,
        (25 / 1280, 153 / 681, 1229 / 1280, 618 / 681))
text(s, 'Capture de démo : la partie affichée est en mode local.', 55, 499, 690, 14, 8, MUTED)
notes(s, "30 s — À la fin d’un run, on retrouve le score, les pièces, les dernières secondes de la course et le classement. Les highlights restent locaux. Cette capture fournie affiche explicitement une partie locale : ne pas présenter les 2 930 points comme une transaction confirmée. Le classement et le résultat confirmé sont distingués par l’interface. Accroche : ‘Le score de ton pote devient soudain une affaire personnelle.’")


# 5 — Le rôle exact de Monad, avec une invitation à la démo.
s = base(5, 'MONAD TESTNET  /  PLACE À LA DÉMO')
text(s, 'DU CARDIO.', 52, 83, 622, 63, 47, WHITE, True, True)
text(s, 'ET DU ONCHAIN.', 50, 138, 760, 63, 47, YELLOW, True, True)
steps = [
    (52, '01', 'TU JOUES', 'Webcam + simulation', 'Dans le navigateur', YELLOW),
    (357, '02', 'ON VÉRIFIE', 'Rejeu côté serveur', 'Score + pièces recalculés', YELLOW),
    (662, '03', 'MONAD GARDE', 'Pseudo, scores, pièces', 'Replays + top 25', PURPLE),
]
for x, number, title, first, second, color in steps:
    shape(s, x, 241, 246, 135, PANEL, LINE, 1, True)
    shape(s, x + 18, 258, 32, 24, color, radius=True)
    text(s, number, x + 18, 263, 32, 16, 10, DEEP if color == YELLOW else WHITE, True, align='center')
    text(s, title, x + 18, 293, 213, 27, 18, WHITE, True)
    text(s, first + '\n' + second, x + 18, 327, 219, 41, 13.5, MUTED, spacing=1.4)
for x in [310, 615]:
    shape(s, x, 290, 34, 24, YELLOW, kind=MSO_SHAPE.RIGHT_ARROW)
text(s, 'Pas de wallet à connecter : le relayer paie les transactions.', 54, 394, 854, 27, 17, WHITE)
chip(s, 'NEXT LEVEL', 53, 438, 108, PURPLE, WHITE, 9)
text(s, 'Une tx par mouvement : API prête à raccorder.', 172, 444, 515, 25, 14, MUTED)
shape(s, 704, 434, 204, 43, YELLOW, radius=True)
button = text(s, 'À TOI DE COURIR.', 711, 447, 190, 28, 16, DEEP, True, True, align='center')
button.text_frame.paragraphs[0].runs[0].hyperlink.address = 'https://hackathon-monad-brown.vercel.app'
notes(s, "40 s — Le jeu reste local et fluide. En fin de partie, le serveur rejoue les commandes pour recalculer le score et les pièces, puis le relayer enregistre le résultat sur Monad testnet. Le joueur n’a pas besoin de wallet : nous payons le gas. Le contrat conserve le pseudo, les résultats, les replays et le top 25. Le rejeu vérifie la cohérence, pas la réalité des gestes. Pour la suite, l’API une transaction par changement de commande est prête sur chain, avec son propre contrat et relayer ; son déploiement et son raccordement au front restent à faire. Conclure par une démo : ‘Qui veut prendre le record ?’ Le bouton ouvre https://hackathon-monad-brown.vercel.app.")

output = ROOT / 'Subway-Frauder-Presentation.pptx'
prs.save(output)
print(f'{len(prs.slides)} slides : {output}')
